# -*- coding: utf-8 -*-
"""
DivisionSellGeneralAdminCost_Allocation_Cmd.py

受け取った入力ファイル群（SellGeneralAdminCost_Allocation_Cmd_0002.py と同等）から、
損益計算書TSVを対象に 1列目と「販売費及び一般管理費計」列のみを残した
Div販管費TSVを出力する。

さらに、step0001 の出力TSVを読み込み、
「スタートアップコミュニティDiv販管費」〜「C008_新規プロポーザル」
の2列目合計を算出して、各終端行の直下に「Div販管費」行を挿入した
step0002 TSV を出力する。
"""

from __future__ import annotations

import csv
import os
import re
import sys
from datetime import datetime
from typing import Dict, List, Optional, Tuple

import openpyxl

PL_TSV_PATTERN = re.compile(r"^損益計算書_(\d{4})年(\d{2})月_A∪B_プロジェクト名_C∪D_vertical\.tsv$")
STEP0001_TSV_PATTERN = re.compile(r"^損益計算書_step0001_(\d{4})年(\d{2})月_A∪B_C∪D_Div販管費_vertical\.tsv$")
STEP0002_TSV_PATTERN = re.compile(r"^損益計算書_step0002_(\d{4})年(\d{2})月_A∪B_C∪D_Div販管費_vertical\.tsv$")
STEP0003_TSV_PATTERN = re.compile(r"^損益計算書_step0003_(\d{4})年(\d{2})月_A∪B_C∪D_Div販管費_vertical\.tsv$")
MANHOUR_TSV_PATTERN = re.compile(
    r"^工数_(\d{4})年(\d{2})月_step0014_各プロジェクトの計上カンパニー名_工数_カンパニーの工数\.tsv$"
)
TARGET_COLUMN_NAME: str = "販売費及び一般管理費計"
RANGE_START_LABEL: str = "スタートアップコミュニティDiv販管費"
RANGE_END_LABEL: str = "C008_新規プロポーザル"
PROJECT_CODE_PATTERN = re.compile(r"(P\d{5}|[A-OQ-Z]\d{3})_")


def append_error_log(pszMessage: str) -> None:
    pszOutputPath: str = os.path.join(
        os.path.dirname(__file__),
        "DivisionSellGeneralAdminCost_Allocation_Cmd_error.txt",
    )
    with open(pszOutputPath, "a", encoding="utf-8", newline="") as objFile:
        objFile.write(pszMessage + "\n")


def write_numeric_warning_error_file(
    pszInputPath: str,
    objWarnings: List[Dict[str, str]],
) -> str:
    pszNow: str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    append_error_log(f"[{pszNow}] ERROR START")
    append_error_log("JOB: DivisionSellGeneralAdminCost_Allocation_Cmd.py")
    append_error_log("ERROR_TYPE: NON_NUMERIC_VALUE_REPLACED_WITH_ZERO")
    append_error_log("INPUT_FILE: " + pszInputPath)
    for objWarning in objWarnings:
        append_error_log("ROW_INDEX: " + objWarning["row_index"])
        append_error_log("VALUE: " + objWarning["value"])
        append_error_log("REPLACED_VALUE: 0")
    append_error_log("RESULT: WARNING")
    pszEnd: str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    append_error_log(f"[{pszEnd}] ERROR END")

    pszBaseName: str = os.path.basename(pszInputPath)
    pszErrorPath: str = os.path.join(os.path.dirname(pszInputPath), pszBaseName + "_error.txt")
    objLines: List[str] = [
        "JOB: DivisionSellGeneralAdminCost_Allocation_Cmd.py",
        "ERROR_TYPE: NON_NUMERIC_VALUE_REPLACED_WITH_ZERO",
        "INPUT_FILE: " + pszInputPath,
    ]
    for objWarning in objWarnings:
        objLines.append("ROW_INDEX: " + objWarning["row_index"])
        objLines.append("VALUE: " + objWarning["value"])
        objLines.append("REPLACED_VALUE: 0")
    objLines.append("RESULT: WARNING")
    with open(pszErrorPath, "w", encoding="utf-8", newline="") as objFile:
        objFile.write("\n".join(objLines) + "\n")
    return pszErrorPath


def is_pl_tsv_file_name(pszBaseName: str) -> bool:
    return PL_TSV_PATTERN.fullmatch(pszBaseName) is not None


def is_step0001_tsv_file_name(pszBaseName: str) -> bool:
    return STEP0001_TSV_PATTERN.fullmatch(pszBaseName) is not None


def is_step0002_tsv_file_name(pszBaseName: str) -> bool:
    return STEP0002_TSV_PATTERN.fullmatch(pszBaseName) is not None


def is_manhour_tsv_file_name(pszBaseName: str) -> bool:
    return MANHOUR_TSV_PATTERN.fullmatch(pszBaseName) is not None


def build_step0001_output_file_name(pszInputBaseName: str) -> Optional[str]:
    objMatch = PL_TSV_PATTERN.fullmatch(pszInputBaseName)
    if objMatch is None:
        return None
    iYear: str = objMatch.group(1)
    iMonth: str = objMatch.group(2)
    return f"損益計算書_step0001_{iYear}年{iMonth}月_A∪B_C∪D_Div販管費_vertical.tsv"


def build_step0002_output_file_name(pszStep0001BaseName: str) -> Optional[str]:
    objMatch = STEP0001_TSV_PATTERN.fullmatch(pszStep0001BaseName)
    if objMatch is None:
        return None
    iYear: str = objMatch.group(1)
    iMonth: str = objMatch.group(2)
    return f"損益計算書_step0002_{iYear}年{iMonth}月_A∪B_C∪D_Div販管費_vertical.tsv"


def build_step0003_output_file_name(pszStep0002BaseName: str) -> Optional[str]:
    objMatch = STEP0002_TSV_PATTERN.fullmatch(pszStep0002BaseName)
    if objMatch is None:
        return None
    iYear: str = objMatch.group(1)
    iMonth: str = objMatch.group(2)
    return f"損益計算書_step0003_{iYear}年{iMonth}月_A∪B_C∪D_Div販管費_vertical.tsv"


def build_step0004_output_file_name(pszStep0003BaseName: str) -> Optional[str]:
    objMatch = STEP0003_TSV_PATTERN.fullmatch(pszStep0003BaseName)
    if objMatch is None:
        return None
    iYear: str = objMatch.group(1)
    iMonth: str = objMatch.group(2)
    return f"損益計算書_step0004_{iYear}年{iMonth}月_A∪B_C∪D_Div販管費_vertical.tsv"


def extract_year_month_from_name(pszBaseName: str) -> Optional[str]:
    objPatterns = [PL_TSV_PATTERN, STEP0001_TSV_PATTERN, STEP0002_TSV_PATTERN, MANHOUR_TSV_PATTERN]
    for objPattern in objPatterns:
        objMatch = objPattern.fullmatch(pszBaseName)
        if objMatch is not None:
            return f"{objMatch.group(1)}年{objMatch.group(2)}月"
    return None


def extract_project_code(pszText: str) -> Optional[str]:
    objMatch = PROJECT_CODE_PATTERN.search(pszText)
    if objMatch is None:
        return None
    return objMatch.group(1)


def normalize_project_code_for_match(pszText: str) -> Optional[str]:
    return extract_project_code(pszText)


def find_target_column_index(objHeader: List[str]) -> Optional[int]:
    for iIndex, pszColumn in enumerate(objHeader):
        if pszColumn == TARGET_COLUMN_NAME:
            return iIndex
    return None


def process_one_pl_tsv_to_step0001(pszInputPath: str) -> str:
    pszBaseName: str = os.path.basename(pszInputPath)
    pszOutputBaseName: Optional[str] = build_step0001_output_file_name(pszBaseName)
    if pszOutputBaseName is None:
        raise ValueError(f"invalid PL TSV file name: {pszBaseName}")

    pszOutputPath: str = os.path.join(os.path.dirname(pszInputPath), pszOutputBaseName)

    with open(pszInputPath, "r", encoding="utf-8", newline="") as objInputFile:
        objReader = csv.reader(objInputFile, delimiter="\t")
        objRows: List[List[str]] = list(objReader)

    if not objRows:
        raise ValueError(f"empty TSV: {pszInputPath}")

    objHeader: List[str] = objRows[0]
    if len(objHeader) == 0:
        raise ValueError(f"header is empty: {pszInputPath}")

    iTargetIndex: Optional[int] = find_target_column_index(objHeader)
    if iTargetIndex is None:
        raise ValueError(
            f"required column not found: {TARGET_COLUMN_NAME} in {pszInputPath}"
        )

    iFirstColumnIndex: int = 0
    objOutputRows: List[List[str]] = []
    for objRow in objRows:
        pszFirstValue: str = objRow[iFirstColumnIndex] if len(objRow) > iFirstColumnIndex else ""
        pszTargetValue: str = objRow[iTargetIndex] if len(objRow) > iTargetIndex else ""
        objOutputRows.append([pszFirstValue, pszTargetValue])

    with open(pszOutputPath, "w", encoding="utf-8", newline="") as objOutputFile:
        objWriter = csv.writer(objOutputFile, delimiter="\t", lineterminator="\n")
        objWriter.writerows(objOutputRows)

    return pszOutputPath


def parse_numeric_value(
    pszValue: str,
    iRowIndex: int,
    objWarnings: List[Dict[str, str]],
) -> float:
    pszTrimmed: str = pszValue.strip()
    if pszTrimmed == "":
        return 0.0

    pszNormalized: str = pszTrimmed.replace(",", "")
    try:
        return float(pszNormalized)
    except ValueError:
        objWarnings.append({"row_index": str(iRowIndex), "value": pszValue})
        return 0.0


def process_one_step0001_tsv_to_step0002(pszStep0001Path: str) -> Tuple[Optional[str], Optional[str]]:
    pszBaseName: str = os.path.basename(pszStep0001Path)
    pszOutputBaseName: Optional[str] = build_step0002_output_file_name(pszBaseName)
    if pszOutputBaseName is None:
        raise ValueError(f"invalid step0001 TSV file name: {pszBaseName}")

    with open(pszStep0001Path, "r", encoding="utf-8", newline="") as objInputFile:
        objReader = csv.reader(objInputFile, delimiter="\t")
        objRows: List[List[str]] = list(objReader)

    if not objRows:
        return None, None

    objOutputRows: List[List[str]] = []
    bInRange: bool = False
    fCurrentRangeSum: float = 0.0
    iInsertedCount: int = 0
    objWarnings: List[Dict[str, str]] = []

    for iRowIndex, objRow in enumerate(objRows):
        objOutputRows.append(objRow)
        pszFirstColumn: str = objRow[0] if len(objRow) >= 1 else ""
        pszSecondColumn: str = objRow[1] if len(objRow) >= 2 else ""

        if pszFirstColumn == RANGE_START_LABEL:
            bInRange = True
            fCurrentRangeSum = 0.0

        if bInRange:
            fCurrentRangeSum += parse_numeric_value(pszSecondColumn, iRowIndex, objWarnings)

        if bInRange and pszFirstColumn == RANGE_END_LABEL:
            objOutputRows.append(["Div販管費", str(int(fCurrentRangeSum))])
            iInsertedCount += 1
            bInRange = False
            fCurrentRangeSum = 0.0

    if iInsertedCount == 0:
        return None, None

    pszOutputPath: str = os.path.join(os.path.dirname(pszStep0001Path), pszOutputBaseName)
    with open(pszOutputPath, "w", encoding="utf-8", newline="") as objOutputFile:
        objWriter = csv.writer(objOutputFile, delimiter="\t", lineterminator="\n")
        objWriter.writerows(objOutputRows)

    pszErrorPath: Optional[str] = None
    if objWarnings:
        pszErrorPath = write_numeric_warning_error_file(pszStep0001Path, objWarnings)

    return pszOutputPath, pszErrorPath


def load_manhour_map(pszManhourPath: str) -> Dict[str, str]:
    with open(pszManhourPath, "r", encoding="utf-8", newline="") as objInputFile:
        objReader = csv.reader(objInputFile, delimiter="\t")
        objRows: List[List[str]] = list(objReader)

    objMap: Dict[str, str] = {}
    for objRow in objRows:
        if len(objRow) < 3:
            continue
        pszProjectCode: Optional[str] = extract_project_code(objRow[0])
        if pszProjectCode is None:
            continue
        objMap[pszProjectCode] = objRow[2]
    return objMap


def process_one_step0002_with_manhour_to_step0003(
    pszStep0002Path: str,
    pszManhourPath: str,
) -> str:
    pszStep0002BaseName: str = os.path.basename(pszStep0002Path)
    pszOutputBaseName: Optional[str] = build_step0003_output_file_name(pszStep0002BaseName)
    if pszOutputBaseName is None:
        raise ValueError(f"invalid step0002 TSV file name: {pszStep0002BaseName}")

    objManhourMap: Dict[str, str] = load_manhour_map(pszManhourPath)

    with open(pszStep0002Path, "r", encoding="utf-8", newline="") as objInputFile:
        objReader = csv.reader(objInputFile, delimiter="\t")
        objRows: List[List[str]] = list(objReader)

    objOutputRows: List[List[str]] = []
    for iRowIndex, objRow in enumerate(objRows):
        objNewRow: List[str] = list(objRow)
        while len(objNewRow) < 3:
            objNewRow.append("")
        if iRowIndex == 0:
            objNewRow[2] = "工数"
            objOutputRows.append(objNewRow)
            continue
        pszProjectCode: Optional[str] = extract_project_code(objNewRow[0] if objNewRow else "")
        if pszProjectCode is not None and pszProjectCode in objManhourMap:
            objNewRow[2] = objManhourMap[pszProjectCode]
        elif objNewRow[2] == "":
            objNewRow[2] = "0:00:00"
        objOutputRows.append(objNewRow)

    pszOutputPath: str = os.path.join(os.path.dirname(pszStep0002Path), pszOutputBaseName)
    with open(pszOutputPath, "w", encoding="utf-8", newline="") as objOutputFile:
        objWriter = csv.writer(objOutputFile, delimiter="\t", lineterminator="\n")
        objWriter.writerows(objOutputRows)
    return pszOutputPath


def parse_time_to_seconds(pszTimeText: str) -> float:
    pszText: str = pszTimeText.strip()
    if pszText == "":
        return 0.0
    objParts: List[str] = pszText.split(":")
    if len(objParts) != 3:
        return 0.0
    try:
        iHours: int = int(objParts[0])
        iMinutes: int = int(objParts[1])
        iSeconds: int = int(objParts[2])
    except ValueError:
        return 0.0
    return float(iHours * 3600 + iMinutes * 60 + iSeconds)


def format_seconds_to_time_text(fSeconds: float) -> str:
    iTotalSeconds: int = int(round(fSeconds))
    if iTotalSeconds < 0:
        iTotalSeconds = 0
    iHours: int = iTotalSeconds // 3600
    iRemainSeconds: int = iTotalSeconds % 3600
    iMinutes: int = iRemainSeconds // 60
    iSeconds: int = iRemainSeconds % 60
    return f"{iHours}:{iMinutes:02d}:{iSeconds:02d}"


def process_one_step0003_to_step0004(pszStep0003Path: str) -> Optional[str]:
    pszBaseName: str = os.path.basename(pszStep0003Path)
    pszOutputBaseName: Optional[str] = build_step0004_output_file_name(pszBaseName)
    if pszOutputBaseName is None:
        return None
    with open(pszStep0003Path, "r", encoding="utf-8", newline="") as objInputFile:
        objReader = csv.reader(objInputFile, delimiter="\t")
        objRows: List[List[str]] = list(objReader)
    if not objRows:
        return None

    iDivRowIndex: int = -1
    for iRowIndex, objRow in enumerate(objRows):
        if len(objRow) >= 1 and objRow[0].strip() == "Div販管費":
            iDivRowIndex = iRowIndex
            break
    if iDivRowIndex < 0:
        return None

    fTotalAmount: float = 0.0
    if len(objRows[iDivRowIndex]) >= 2:
        fTotalAmount = parse_numeric_value(objRows[iDivRowIndex][1], iDivRowIndex, [])
    iTargetTotal: int = int(round(fTotalAmount))

    objTargetIndices: List[int] = []
    objSeconds: List[float] = []
    fTotalSeconds: float = 0.0
    for iRowIndex in range(iDivRowIndex + 1, len(objRows)):
        objRow = objRows[iRowIndex]
        while len(objRow) < 3:
            objRow.append("")
        fSeconds: float = parse_time_to_seconds(objRow[2])
        objTargetIndices.append(iRowIndex)
        objSeconds.append(fSeconds)
        if fSeconds > 0:
            fTotalSeconds += fSeconds

    if not objTargetIndices:
        return None

    objRawValues: List[float] = []
    for fSeconds in objSeconds:
        if fTotalSeconds <= 0.0 or fSeconds <= 0.0:
            objRawValues.append(0.0)
        else:
            objRawValues.append(float(iTargetTotal) * fSeconds / fTotalSeconds)
    objBaseValues: List[int] = [int(fRawValue // 1) for fRawValue in objRawValues]
    iRemain: int = iTargetTotal - sum(objBaseValues)
    objRankIndices: List[int] = list(range(len(objTargetIndices)))
    objRankIndices.sort(
        key=lambda iIndex: (
            objRawValues[iIndex] - objBaseValues[iIndex],
            objSeconds[iIndex],
            -objTargetIndices[iIndex],
        ),
        reverse=True,
    )
    if iRemain > 0:
        for iIndex in objRankIndices[:iRemain]:
            objBaseValues[iIndex] += 1
    elif iRemain < 0:
        objRankIndicesAsc: List[int] = list(reversed(objRankIndices))
        for iIndex in objRankIndicesAsc[: (-iRemain)]:
            objBaseValues[iIndex] -= 1

    for iTargetIndex, iRowIndex in enumerate(objTargetIndices):
        objRows[iRowIndex][1] = str(objBaseValues[iTargetIndex])

    pszOutputPath: str = os.path.join(os.path.dirname(pszStep0003Path), pszOutputBaseName)
    with open(pszOutputPath, "w", encoding="utf-8", newline="") as objOutputFile:
        objWriter = csv.writer(objOutputFile, delimiter="\t", lineterminator="\n")
        objWriter.writerows(objRows)
    return pszOutputPath


def build_year_month_path_map(objPaths: List[str]) -> Dict[str, str]:
    objMap: Dict[str, str] = {}
    for pszPath in objPaths:
        pszYearMonth: Optional[str] = extract_year_month_from_name(os.path.basename(pszPath))
        if pszYearMonth is None:
            continue
        objMap[pszYearMonth] = pszPath
    return objMap


def run_step0003_with_maps(
    objStep0002Map: Dict[str, str],
    objManhourMap: Dict[str, str],
) -> Tuple[List[str], List[str]]:
    objStep0003OutputPaths: List[str] = []
    objSkippedYearMonths: List[str] = []
    for pszYearMonth, pszStep0002Path in sorted(objStep0002Map.items()):
        if pszYearMonth not in objManhourMap:
            objSkippedYearMonths.append(pszYearMonth)
            continue
        pszStep0003Path: str = process_one_step0002_with_manhour_to_step0003(
            pszStep0002Path,
            objManhourMap[pszYearMonth],
        )
        objStep0003OutputPaths.append(pszStep0003Path)
    return objStep0003OutputPaths, objSkippedYearMonths


def run_step0004_for_paths(objStep0003Paths: List[str]) -> List[str]:
    objStep0004Paths: List[str] = []
    for pszStep0003Path in objStep0003Paths:
        pszStep0004Path: Optional[str] = process_one_step0003_to_step0004(pszStep0003Path)
        if pszStep0004Path is not None:
            objStep0004Paths.append(pszStep0004Path)
    return objStep0004Paths


def parse_year_month_value(pszYearMonth: str) -> Optional[Tuple[int, int]]:
    try:
        iYearText: str = pszYearMonth.split("年", 1)[0]
        iMonthText: str = pszYearMonth.split("年", 1)[1].split("月", 1)[0]
        iYear: int = int(iYearText)
        iMonth: int = int(iMonthText)
    except (ValueError, IndexError):
        return None
    if iMonth < 1 or iMonth > 12:
        return None
    return iYear, iMonth


def parse_account_periods_from_file(pszPath: str) -> Dict[str, Tuple[Tuple[int, int], Tuple[int, int]]]:
    with open(pszPath, "r", encoding="utf-8", newline="") as objInputFile:
        objLines: List[str] = [pszLine.rstrip("\n").rstrip("\r") for pszLine in objInputFile]
    objPeriods: Dict[str, Tuple[Tuple[int, int], Tuple[int, int]]] = {}
    pszCurrentSection: str = ""
    pszCurrentPeriod: str = ""
    for iIndex, pszLine in enumerate(objLines):
        pszText: str = pszLine.strip()
        if pszText in ("3月決算の会計期間:", "8月決算の会計期間:"):
            pszCurrentSection = pszText[:3]
            pszCurrentPeriod = ""
            continue
        if pszText in ("前期", "当期"):
            pszCurrentPeriod = pszText
            continue
        if not pszText.startswith("開始:"):
            continue
        if pszCurrentSection == "" or pszCurrentPeriod == "":
            continue
        if iIndex + 1 >= len(objLines):
            continue
        pszStart = pszText.split(":", 1)[1].strip().replace("/", "年") + "月"
        pszEndLine: str = objLines[iIndex + 1].strip()
        if not pszEndLine.startswith("終了:"):
            continue
        pszEnd = pszEndLine.split(":", 1)[1].strip().replace("/", "年") + "月"
        objStart = parse_year_month_value(pszStart)
        objEnd = parse_year_month_value(pszEnd)
        if objStart is None or objEnd is None:
            continue
        objPeriods[f"{pszCurrentSection}_{pszCurrentPeriod}"] = (objStart, objEnd)
    return objPeriods


def is_within_period(objTarget: Tuple[int, int], objStart: Tuple[int, int], objEnd: Tuple[int, int]) -> bool:
    return objStart <= objTarget <= objEnd


def aggregate_step0004_for_periods(objStep0004Paths: List[str]) -> Tuple[List[str], List[str], List[str]]:
    pszPeriodFilePath: str = os.path.join(
        os.path.dirname(__file__),
        "SellGeneralAdminCost_Allocation_Cmd_SelectedRange_And_AccountPeriodRange.txt",
    )
    if not os.path.isfile(pszPeriodFilePath):
        return [], [], []
    objPeriods = parse_account_periods_from_file(pszPeriodFilePath)
    if not objPeriods:
        return [], [], []

    objStep0004ByMonth: Dict[Tuple[int, int], str] = {}
    for pszPath in objStep0004Paths:
        objMatch = re.fullmatch(r".*step0004_(\d{4})年(\d{2})月_.*\.tsv$", os.path.basename(pszPath))
        if objMatch is None:
            continue
        objStep0004ByMonth[(int(objMatch.group(1)), int(objMatch.group(2)))] = pszPath

    objOutputPaths: List[str] = []
    objWarningPaths: List[str] = []
    objErrorPaths: List[str] = []
    for pszPeriodKey, (objStart, objEnd) in objPeriods.items():
        objMonths: List[Tuple[int, int]] = []
        iYear, iMonth = objStart
        while (iYear, iMonth) <= objEnd:
            objMonths.append((iYear, iMonth))
            iMonth += 1
            if iMonth == 13:
                iMonth = 1
                iYear += 1
        objMissingMonths: List[str] = []
        objSourcePaths: List[str] = []
        for objMonth in objMonths:
            if objMonth not in objStep0004ByMonth:
                objMissingMonths.append(f"{objMonth[0]}年{objMonth[1]:02d}月")
            else:
                objSourcePaths.append(objStep0004ByMonth[objMonth])
        if objMissingMonths:
            pszErrorPath = os.path.join(
                os.path.dirname(objStep0004Paths[0]),
                f"損益計算書_step0004_{objStart[0]}年{objStart[1]:02d}月-{objEnd[0]}年{objEnd[1]:02d}月_A∪B_C∪D_Div販管費_vertical.tsv_error.txt",
            )
            with open(pszErrorPath, "w", encoding="utf-8", newline="") as objErrorFile:
                objErrorFile.write("ERROR_TYPE: MISSING_MONTH_IN_PERIOD\n")
                objErrorFile.write("PERIOD: " + pszPeriodKey + "\n")
                for pszMonth in objMissingMonths:
                    objErrorFile.write("MISSING_MONTH: " + pszMonth + "\n")
            objErrorPaths.append(pszErrorPath)
            continue
        if not objSourcePaths:
            continue
        objAggregatedRows: List[List[str]] = []
        objNumericWarnings: List[str] = []
        for iFileIndex, pszSourcePath in enumerate(objSourcePaths):
            with open(pszSourcePath, "r", encoding="utf-8", newline="") as objInputFile:
                objRows = list(csv.reader(objInputFile, delimiter="\t"))
            if iFileIndex == 0:
                objAggregatedRows = [list(objRow) for objRow in objRows]
                continue
            for iRowIndex, objRow in enumerate(objRows):
                while iRowIndex >= len(objAggregatedRows):
                    objAggregatedRows.append([""])
                objTargetRow = objAggregatedRows[iRowIndex]
                iMaxColumns = max(len(objTargetRow), len(objRow))
                while len(objTargetRow) < iMaxColumns:
                    objTargetRow.append("")
                while len(objRow) < iMaxColumns:
                    objRow.append("")
                if iRowIndex == 0:
                    objAggregatedRows[iRowIndex] = objTargetRow
                    continue
                for iColumnIndex in range(1, iMaxColumns):
                    if iColumnIndex == 2:
                        pszLeftTime: str = objTargetRow[iColumnIndex].strip()
                        pszRightTime: str = objRow[iColumnIndex].strip()
                        if pszLeftTime != "" and re.fullmatch(r"\d+:\d{2}:\d{2}", pszLeftTime) is None:
                            objNumericWarnings.append(f"ROW={iRowIndex},COL={iColumnIndex},VALUE={objTargetRow[iColumnIndex]}")
                        if pszRightTime != "" and re.fullmatch(r"\d+:\d{2}:\d{2}", pszRightTime) is None:
                            objNumericWarnings.append(f"ROW={iRowIndex},COL={iColumnIndex},VALUE={objRow[iColumnIndex]}")
                        fLeftSeconds: float = parse_time_to_seconds(objTargetRow[iColumnIndex])
                        fRightSeconds: float = parse_time_to_seconds(objRow[iColumnIndex])
                        objTargetRow[iColumnIndex] = format_seconds_to_time_text(fLeftSeconds + fRightSeconds)
                        continue
                    fLeft = parse_numeric_value(objTargetRow[iColumnIndex], iRowIndex, [])
                    fRightWarnings: List[Dict[str, str]] = []
                    fRight = parse_numeric_value(objRow[iColumnIndex], iRowIndex, fRightWarnings)
                    if fRightWarnings:
                        objNumericWarnings.append(f"ROW={iRowIndex},COL={iColumnIndex},VALUE={objRow[iColumnIndex]}")
                    objTargetRow[iColumnIndex] = str(int(round(fLeft + fRight)))
                objAggregatedRows[iRowIndex] = objTargetRow
        pszOutputPath = os.path.join(
            os.path.dirname(objSourcePaths[0]),
            f"損益計算書_step0004_{objStart[0]}年{objStart[1]:02d}月-{objEnd[0]}年{objEnd[1]:02d}月_A∪B_C∪D_Div販管費_vertical.tsv",
        )
        with open(pszOutputPath, "w", encoding="utf-8", newline="") as objOutputFile:
            objWriter = csv.writer(objOutputFile, delimiter="\t", lineterminator="\n")
            objWriter.writerows(objAggregatedRows)
        objOutputPaths.append(pszOutputPath)
        if objNumericWarnings:
            pszWarningPath = pszOutputPath + "_error.txt"
            with open(pszWarningPath, "w", encoding="utf-8", newline="") as objWarningFile:
                objWarningFile.write("ERROR_TYPE: NON_NUMERIC_VALUE_REPLACED_WITH_ZERO\n")
                objWarningFile.write("PERIOD: " + pszPeriodKey + "\n")
                for pszWarning in objNumericWarnings:
                    objWarningFile.write(pszWarning + "\n")
            objWarningPaths.append(pszWarningPath)
    return objOutputPaths, objWarningPaths, objErrorPaths


def create_step0004_summary_excel(objPeriodPaths: List[str]) -> Optional[str]:
    if not objPeriodPaths:
        return None
    objWorkbook = openpyxl.Workbook()
    objDefaultSheet = objWorkbook.active
    objWorkbook.remove(objDefaultSheet)

    objPattern = re.compile(r"step0004_(\d{4}年\d{2}月-\d{4}年\d{2}月)_A∪B_C∪D_Div販管費_vertical\.tsv$")
    for pszPeriodPath in objPeriodPaths:
        pszBaseName: str = os.path.basename(pszPeriodPath)
        objMatch = objPattern.search(pszBaseName)
        pszSheetName: str = objMatch.group(1) if objMatch is not None else os.path.splitext(pszBaseName)[0][:31]
        if len(pszSheetName) > 31:
            pszSheetName = pszSheetName[:31]
        objSheet = objWorkbook.create_sheet(title=pszSheetName)
        with open(pszPeriodPath, "r", encoding="utf-8", newline="") as objInputFile:
            objRows: List[List[str]] = list(csv.reader(objInputFile, delimiter="\t"))
        for iRowIndex, objRow in enumerate(objRows, start=1):
            for iColumnIndex, pszValue in enumerate(objRow, start=1):
                objCell = objSheet.cell(row=iRowIndex, column=iColumnIndex)
                if iRowIndex == 1 or iColumnIndex == 1:
                    objCell.value = str(pszValue)
                    continue
                if iColumnIndex == 2:
                    try:
                        objCell.value = int(float(str(pszValue).replace(",", "").strip()))
                    except ValueError:
                        objCell.value = 0
                    continue
                if iColumnIndex == 3:
                    pszText = str(pszValue).strip()
                    objCell.value = pszText if pszText != "" else "0:00:00"
                    continue
                objCell.value = str(pszValue)
        objSheet.column_dimensions["A"].width = 105
        objSheet.column_dimensions["B"].width = 28
        objSheet.column_dimensions["C"].width = 10
    pszOutputPath: str = os.path.join(
        os.path.dirname(objPeriodPaths[0]),
        "損益計算書_step0004_合算Div販管費.xlsx",
    )
    objWorkbook.save(pszOutputPath)
    return pszOutputPath


def build_step0005_output_file_name(pszStep0004PeriodBaseName: str) -> Optional[str]:
    objMatch = re.fullmatch(
        r"損益計算書_step0004_(\d{4}年\d{2}月-\d{4}年\d{2}月)_A∪B_C∪D_Div販管費_vertical\.tsv",
        pszStep0004PeriodBaseName,
    )
    if objMatch is None:
        return None
    return f"損益計算書_step0005_{objMatch.group(1)}_A∪B_C∪D_Div販管費_vertical.tsv"


def create_step0005_from_period_paths(objPeriodPaths: List[str]) -> List[str]:
    objStep0005Paths: List[str] = []
    for pszPeriodPath in objPeriodPaths:
        pszBaseName: str = os.path.basename(pszPeriodPath)
        pszStep0005BaseName: Optional[str] = build_step0005_output_file_name(pszBaseName)
        if pszStep0005BaseName is None:
            continue
        with open(pszPeriodPath, "r", encoding="utf-8", newline="") as objInputFile:
            objRows: List[List[str]] = list(csv.reader(objInputFile, delimiter="\t"))
        iStartIndex: int = -1
        iEndIndex: int = -1
        for iRowIndex, objRow in enumerate(objRows):
            pszFirstColumn: str = objRow[0].strip() if len(objRow) >= 1 else ""
            if iStartIndex < 0 and pszFirstColumn == "合計":
                iStartIndex = iRowIndex
                continue
            if iStartIndex >= 0 and pszFirstColumn == "Div販管費":
                iEndIndex = iRowIndex
                break
        if iStartIndex >= 0 and iEndIndex >= iStartIndex:
            objRows = objRows[:iStartIndex] + objRows[iEndIndex + 1 :]
        pszOutputPath: str = os.path.join(os.path.dirname(pszPeriodPath), pszStep0005BaseName)
        with open(pszOutputPath, "w", encoding="utf-8", newline="") as objOutputFile:
            objWriter = csv.writer(objOutputFile, delimiter="\t", lineterminator="\n")
            objWriter.writerows(objRows)
        objStep0005Paths.append(pszOutputPath)
    return objStep0005Paths


def load_jurisdiction_master() -> Dict[str, Tuple[str, str]]:
    pszBaseDirectory: str = os.path.dirname(__file__)
    pszTsvPath: str = os.path.join(pszBaseDirectory, "管轄PJ表.tsv")
    pszCsvPath: str = os.path.join(pszBaseDirectory, "管轄PJ表.csv")
    pszInputPath: str = pszTsvPath if os.path.isfile(pszTsvPath) else pszCsvPath
    if not os.path.isfile(pszInputPath):
        return {}

    pszDelimiter: str = "\t" if pszInputPath.lower().endswith(".tsv") else ","
    with open(pszInputPath, "r", encoding="utf-8", newline="") as objInputFile:
        objRows: List[List[str]] = list(csv.reader(objInputFile, delimiter=pszDelimiter))
    if not objRows:
        return {}
    objHeader: List[str] = objRows[0]
    try:
        iCodeIndex: int = objHeader.index("PJコード")
        iCompanyIndex: int = objHeader.index("計上カンパニー")
        iGroupIndex: int = objHeader.index("計上グループ")
    except ValueError:
        return {}

    objMaster: Dict[str, Tuple[str, str]] = {}
    for objRow in objRows[1:]:
        if iCodeIndex >= len(objRow):
            continue
        pszCodeText: str = objRow[iCodeIndex].strip()
        pszCode: Optional[str] = normalize_project_code_for_match(pszCodeText)
        if pszCode is None or pszCode in objMaster:
            continue
        pszCompany: str = objRow[iCompanyIndex].strip() if iCompanyIndex < len(objRow) else ""
        pszGroup: str = objRow[iGroupIndex].strip() if iGroupIndex < len(objRow) else ""
        objMaster[pszCode] = (pszCompany, pszGroup)
    return objMaster


def build_step0006_output_file_name(pszStep0005BaseName: str, pszSuffix: str) -> Optional[str]:
    objMatch = re.fullmatch(
        r"損益計算書_step0005_(\d{4}年\d{2}月-\d{4}年\d{2}月)_A∪B_C∪D_Div販管費_vertical\.tsv",
        pszStep0005BaseName,
    )
    if objMatch is None:
        return None
    return f"損益計算書_step0006_{objMatch.group(1)}_A∪B_C∪D_Div販管費_{pszSuffix}_vertical.tsv"


def create_step0006_from_step0005_paths(objStep0005Paths: List[str]) -> Tuple[List[str], List[str]]:
    objMaster: Dict[str, Tuple[str, str]] = load_jurisdiction_master()
    if not objMaster:
        return [], []

    objDivPaths: List[str] = []
    objGrpPaths: List[str] = []
    for pszStep0005Path in objStep0005Paths:
        with open(pszStep0005Path, "r", encoding="utf-8", newline="") as objInputFile:
            objRows: List[List[str]] = list(csv.reader(objInputFile, delimiter="\t"))

        objDivRows: List[List[str]] = [list(objRow) for objRow in objRows]
        objGrpRows: List[List[str]] = [list(objRow) for objRow in objRows]
        for iRowIndex in range(1, len(objRows)):
            pszFirstColumn: str = objRows[iRowIndex][0] if len(objRows[iRowIndex]) >= 1 else ""
            pszCode: Optional[str] = normalize_project_code_for_match(pszFirstColumn)
            if pszCode is None or pszCode not in objMaster:
                continue
            pszCompany, pszGroup = objMaster[pszCode]
            if len(objDivRows[iRowIndex]) == 0:
                objDivRows[iRowIndex].append("")
            if len(objGrpRows[iRowIndex]) == 0:
                objGrpRows[iRowIndex].append("")
            objDivRows[iRowIndex][0] = pszCompany
            objGrpRows[iRowIndex][0] = pszGroup

        pszBaseName: str = os.path.basename(pszStep0005Path)
        pszDivBaseName: Optional[str] = build_step0006_output_file_name(pszBaseName, "Div")
        pszGrpBaseName: Optional[str] = build_step0006_output_file_name(pszBaseName, "Grp")
        if pszDivBaseName is None or pszGrpBaseName is None:
            continue

        objTargetSubjects: List[str] = [
            "テクノロジーインキュベーション",
            "コンテンツビジネス",
            "スタートアップサイド",
            "スタートアップコミュニティ",
            "スタートアップグロース",
            "経営管理",
            "事業開発",
            "子会社",
            "投資先",
            "本部",
        ]
        objTargetSet = set(objTargetSubjects)
        objAggregatedRows: List[List[str]] = []
        objSubjectToIndex: Dict[str, int] = {}
        if objDivRows:
            objHeaderRow: List[str] = list(objDivRows[0])
            while len(objHeaderRow) < 3:
                objHeaderRow.append("")
            objAggregatedRows.append(objHeaderRow)
        for iRowIndex in range(1, len(objDivRows)):
            objRow: List[str] = list(objDivRows[iRowIndex])
            while len(objRow) < 3:
                objRow.append("")
            pszSubject: str = objRow[0].strip()
            if pszSubject == "その他":
                continue
            if pszSubject in objTargetSet:
                if pszSubject not in objSubjectToIndex:
                    objSubjectToIndex[pszSubject] = len(objAggregatedRows)
                    objAggregatedRows.append([pszSubject, "0", "0:00:00"])
                iTargetIndex: int = objSubjectToIndex[pszSubject]
                iAmount: int = int(round(parse_numeric_value(objRow[1], iRowIndex, [])))
                iExistingAmount: int = int(round(parse_numeric_value(objAggregatedRows[iTargetIndex][1], iTargetIndex, [])))
                fSeconds: float = parse_time_to_seconds(objRow[2])
                fExistingSeconds: float = parse_time_to_seconds(objAggregatedRows[iTargetIndex][2])
                objAggregatedRows[iTargetIndex][1] = str(iExistingAmount + iAmount)
                objAggregatedRows[iTargetIndex][2] = format_seconds_to_time_text(fExistingSeconds + fSeconds)
            else:
                objAggregatedRows.append(objRow)

        pszDivOutputPath: str = os.path.join(os.path.dirname(pszStep0005Path), pszDivBaseName)
        pszGrpOutputPath: str = os.path.join(os.path.dirname(pszStep0005Path), pszGrpBaseName)
        with open(pszDivOutputPath, "w", encoding="utf-8", newline="") as objOutputFile:
            csv.writer(objOutputFile, delimiter="\t", lineterminator="\n").writerows(objAggregatedRows)
        with open(pszGrpOutputPath, "w", encoding="utf-8", newline="") as objOutputFile:
            csv.writer(objOutputFile, delimiter="\t", lineterminator="\n").writerows(objGrpRows)
        objDivPaths.append(pszDivOutputPath)
        objGrpPaths.append(pszGrpOutputPath)
    return objDivPaths, objGrpPaths


def build_step0007_output_file_name(pszStep0006DivBaseName: str) -> Optional[str]:
    objMatch = re.fullmatch(
        r"損益計算書_step0006_(\d{4}年\d{2}月-\d{4}年\d{2}月)_A∪B_C∪D_Div販管費_Div_vertical\.tsv",
        pszStep0006DivBaseName,
    )
    if objMatch is None:
        return None
    return f"損益計算書_step0007_{objMatch.group(1)}_A∪B_C∪D_Div販管費_Div_vertical.tsv"


def build_step0007_grp_output_file_name(pszStep0006GrpBaseName: str) -> Optional[str]:
    objMatch = re.fullmatch(
        r"損益計算書_step0006_(\d{4}年\d{2}月-\d{4}年\d{2}月)_A∪B_C∪D_Div販管費_Grp_vertical\.tsv",
        pszStep0006GrpBaseName,
    )
    if objMatch is None:
        return None
    return f"損益計算書_step0007_{objMatch.group(1)}_A∪B_C∪D_Div販管費_Grp_vertical.tsv"


def create_step0007_from_step0006_div_paths(objStep0006DivPaths: List[str]) -> List[str]:
    objOutputPaths: List[str] = []
    objTargetSubjects: List[str] = [
        "テクノロジーインキュベーション",
        "コンテンツビジネス",
        "スタートアップサイド",
        "スタートアップコミュニティ",
        "スタートアップグロース",
        "経営管理",
        "事業開発",
        "子会社",
        "投資先",
        "本部",
    ]
    objTargetSet = set(objTargetSubjects)
    for pszStep0006DivPath in objStep0006DivPaths:
        with open(pszStep0006DivPath, "r", encoding="utf-8", newline="") as objInputFile:
            objRows: List[List[str]] = list(csv.reader(objInputFile, delimiter="\t"))
        if not objRows:
            continue

        objOutputRows: List[List[str]] = []
        objHeaderRow: List[str] = list(objRows[0])
        while len(objHeaderRow) < 3:
            objHeaderRow.append("")
        objOutputRows.append(objHeaderRow)

        objSubjectToIndex: Dict[str, int] = {}
        for iRowIndex in range(1, len(objRows)):
            objRow: List[str] = list(objRows[iRowIndex])
            while len(objRow) < 3:
                objRow.append("")
            pszSubject: str = objRow[0].strip()
            if pszSubject == "その他":
                continue
            if pszSubject in objTargetSet:
                if pszSubject not in objSubjectToIndex:
                    objSubjectToIndex[pszSubject] = len(objOutputRows)
                    objOutputRows.append([pszSubject, "0", "0:00:00"])
                iTargetRowIndex: int = objSubjectToIndex[pszSubject]
                iAmount: int = int(round(parse_numeric_value(objRow[1], iRowIndex, [])))
                iCurrentAmount: int = int(round(parse_numeric_value(objOutputRows[iTargetRowIndex][1], iTargetRowIndex, [])))
                fSeconds: float = parse_time_to_seconds(objRow[2])
                fCurrentSeconds: float = parse_time_to_seconds(objOutputRows[iTargetRowIndex][2])
                objOutputRows[iTargetRowIndex][1] = str(iCurrentAmount + iAmount)
                objOutputRows[iTargetRowIndex][2] = format_seconds_to_time_text(fCurrentSeconds + fSeconds)
            else:
                if objRow[2].strip() == "":
                    objRow[2] = "0:00:00"
                elif parse_time_to_seconds(objRow[2]) == 0.0 and objRow[2].strip() != "0:00:00":
                    objRow[2] = "0:00:00"
                objOutputRows.append(objRow)

        pszBaseName: str = os.path.basename(pszStep0006DivPath)
        pszOutputBaseName: Optional[str] = build_step0007_output_file_name(pszBaseName)
        if pszOutputBaseName is None:
            continue
        pszOutputPath: str = os.path.join(os.path.dirname(pszStep0006DivPath), pszOutputBaseName)
        with open(pszOutputPath, "w", encoding="utf-8", newline="") as objOutputFile:
            csv.writer(objOutputFile, delimiter="\t", lineterminator="\n").writerows(objOutputRows)
        objOutputPaths.append(pszOutputPath)
    return objOutputPaths


def create_step0007_from_step0006_grp_paths(objStep0006GrpPaths: List[str]) -> List[str]:
    objOutputPaths: List[str] = []
    objTargetSubjects: List[str] = [
        "受託事業-その他",
        "受託事業-施設運営",
        "自社-その他",
        "自社-施設運営",
    ]
    objTargetSet = set(objTargetSubjects)
    for pszStep0006GrpPath in objStep0006GrpPaths:
        with open(pszStep0006GrpPath, "r", encoding="utf-8", newline="") as objInputFile:
            objRows: List[List[str]] = list(csv.reader(objInputFile, delimiter="\t"))
        if not objRows:
            continue

        objOutputRows: List[List[str]] = []
        objHeaderRow: List[str] = list(objRows[0])
        while len(objHeaderRow) < 3:
            objHeaderRow.append("")
        objOutputRows.append(objHeaderRow)

        objSubjectToIndex: Dict[str, int] = {}
        for iRowIndex in range(1, len(objRows)):
            objRow: List[str] = list(objRows[iRowIndex])
            while len(objRow) < 3:
                objRow.append("")
            pszSubject: str = objRow[0].strip()
            if pszSubject == "その他":
                continue
            if pszSubject in objTargetSet:
                if pszSubject not in objSubjectToIndex:
                    objSubjectToIndex[pszSubject] = len(objOutputRows)
                    objOutputRows.append([pszSubject, "0", "0:00:00"])
                iTargetRowIndex: int = objSubjectToIndex[pszSubject]
                iAmount: int = int(round(parse_numeric_value(objRow[1], iRowIndex, [])))
                iCurrentAmount: int = int(round(parse_numeric_value(objOutputRows[iTargetRowIndex][1], iTargetRowIndex, [])))
                fSeconds: float = parse_time_to_seconds(objRow[2])
                fCurrentSeconds: float = parse_time_to_seconds(objOutputRows[iTargetRowIndex][2])
                objOutputRows[iTargetRowIndex][1] = str(iCurrentAmount + iAmount)
                objOutputRows[iTargetRowIndex][2] = format_seconds_to_time_text(fCurrentSeconds + fSeconds)
            else:
                if objRow[2].strip() == "":
                    objRow[2] = "0:00:00"
                elif parse_time_to_seconds(objRow[2]) == 0.0 and objRow[2].strip() != "0:00:00":
                    objRow[2] = "0:00:00"
                objOutputRows.append(objRow)

        pszBaseName: str = os.path.basename(pszStep0006GrpPath)
        pszOutputBaseName: Optional[str] = build_step0007_grp_output_file_name(pszBaseName)
        if pszOutputBaseName is None:
            continue
        pszOutputPath: str = os.path.join(os.path.dirname(pszStep0006GrpPath), pszOutputBaseName)
        with open(pszOutputPath, "w", encoding="utf-8", newline="") as objOutputFile:
            csv.writer(objOutputFile, delimiter="\t", lineterminator="\n").writerows(objOutputRows)
        objOutputPaths.append(pszOutputPath)
    return objOutputPaths


def create_step0007_div_summary_excel(objStep0007DivPaths: List[str]) -> Tuple[Optional[str], Optional[str]]:
    if not objStep0007DivPaths:
        return None, None
    objWorkbook = openpyxl.Workbook()
    objDefaultSheet = objWorkbook.active
    objWorkbook.remove(objDefaultSheet)
    objPattern = re.compile(
        r"損益計算書_step0007_(\d{4}年\d{2}月-\d{4}年\d{2}月)_A∪B_C∪D_Div販管費_Div_vertical\.tsv$"
    )
    for pszStep0007Path in objStep0007DivPaths:
        pszBaseName: str = os.path.basename(pszStep0007Path)
        objMatch = objPattern.fullmatch(pszBaseName)
        pszSheetName: str = objMatch.group(1) + "_Div" if objMatch is not None else os.path.splitext(pszBaseName)[0][:31]
        if len(pszSheetName) > 31:
            pszSheetName = pszSheetName[:31]
        objSheet = objWorkbook.create_sheet(title=pszSheetName)
        with open(pszStep0007Path, "r", encoding="utf-8", newline="") as objInputFile:
            objRows: List[List[str]] = list(csv.reader(objInputFile, delimiter="\t"))
        for iRowIndex, objRow in enumerate(objRows, start=1):
            for iColumnIndex, pszValue in enumerate(objRow, start=1):
                objSheet.cell(row=iRowIndex, column=iColumnIndex).value = str(pszValue)

    pszOutputPath: str = os.path.join(
        os.path.dirname(objStep0007DivPaths[0]),
        "損益計算書_step0008_A∪B_C∪D_Div販管費_Div.xlsx",
    )
    objWorkbook.save(pszOutputPath)
    pszErrorPath: Optional[str] = None
    if len(objStep0007DivPaths) < 4:
        pszErrorPath = pszOutputPath + "_error.txt"
        with open(pszErrorPath, "w", encoding="utf-8", newline="") as objErrorFile:
            objErrorFile.write("JOB: DivisionSellGeneralAdminCost_Allocation_Cmd.py\n")
            objErrorFile.write("ERROR_TYPE: STEP0007_DIV_INPUT_MISSING\n")
            objErrorFile.write(f"EXPECTED_COUNT: 4\n")
            objErrorFile.write(f"ACTUAL_COUNT: {len(objStep0007DivPaths)}\n")
            objErrorFile.write("RESULT: WARNING\n")
            for pszPath in objStep0007DivPaths:
                objErrorFile.write("INPUT_FILE: " + pszPath + "\n")
    return pszOutputPath, pszErrorPath


def create_step0007_grp_summary_excel(objStep0007GrpPaths: List[str]) -> Tuple[Optional[str], Optional[str]]:
    if not objStep0007GrpPaths:
        return None, None
    objWorkbook = openpyxl.Workbook()
    objDefaultSheet = objWorkbook.active
    objWorkbook.remove(objDefaultSheet)
    objPattern = re.compile(
        r"損益計算書_step0007_(\d{4}年\d{2}月-\d{4}年\d{2}月)_A∪B_C∪D_Div販管費_Grp_vertical\.tsv$"
    )
    for pszStep0007Path in objStep0007GrpPaths:
        pszBaseName: str = os.path.basename(pszStep0007Path)
        objMatch = objPattern.fullmatch(pszBaseName)
        pszSheetName: str = objMatch.group(1) + "_Grp" if objMatch is not None else os.path.splitext(pszBaseName)[0][:31]
        if len(pszSheetName) > 31:
            pszSheetName = pszSheetName[:31]
        objSheet = objWorkbook.create_sheet(title=pszSheetName)
        with open(pszStep0007Path, "r", encoding="utf-8", newline="") as objInputFile:
            objRows: List[List[str]] = list(csv.reader(objInputFile, delimiter="\t"))
        for iRowIndex, objRow in enumerate(objRows, start=1):
            for iColumnIndex, pszValue in enumerate(objRow, start=1):
                objSheet.cell(row=iRowIndex, column=iColumnIndex).value = str(pszValue)

    pszOutputPath: str = os.path.join(
        os.path.dirname(objStep0007GrpPaths[0]),
        "損益計算書_step0008_A∪B_C∪D_Div販管費_Grp.xlsx",
    )
    objWorkbook.save(pszOutputPath)
    pszErrorPath: Optional[str] = None
    if len(objStep0007GrpPaths) < 4:
        pszErrorPath = pszOutputPath + "_error.txt"
        with open(pszErrorPath, "w", encoding="utf-8", newline="") as objErrorFile:
            objErrorFile.write("JOB: DivisionSellGeneralAdminCost_Allocation_Cmd.py\n")
            objErrorFile.write("ERROR_TYPE: STEP0007_GRP_INPUT_MISSING\n")
            objErrorFile.write("EXPECTED_COUNT: 4\n")
            objErrorFile.write(f"ACTUAL_COUNT: {len(objStep0007GrpPaths)}\n")
            objErrorFile.write("RESULT: WARNING\n")
            for pszPath in objStep0007GrpPaths:
                objErrorFile.write("INPUT_FILE: " + pszPath + "\n")
    return pszOutputPath, pszErrorPath


def build_step0008_div_output_file_name(pszStep0007DivBaseName: str) -> Optional[str]:
    objMatch = re.fullmatch(
        r"損益計算書_step0007_(\d{4}年\d{2}月-\d{4}年\d{2}月)_A∪B_C∪D_Div販管費_Div_vertical\.tsv",
        pszStep0007DivBaseName,
    )
    if objMatch is None:
        return None
    return f"損益計算書_step0008_{objMatch.group(1)}_A∪B_C∪D_Div販管費_Div_vertical.tsv"


def build_step0008_grp_output_file_name(pszStep0007GrpBaseName: str) -> Optional[str]:
    objMatch = re.fullmatch(
        r"損益計算書_step0007_(\d{4}年\d{2}月-\d{4}年\d{2}月)_A∪B_C∪D_Div販管費_Grp_vertical\.tsv",
        pszStep0007GrpBaseName,
    )
    if objMatch is None:
        return None
    return f"損益計算書_step0008_{objMatch.group(1)}_A∪B_C∪D_Div販管費_Grp_vertical.tsv"


def create_step0008_from_step0007_div_paths(objStep0007DivPaths: List[str]) -> List[str]:
    objOutputPaths: List[str] = []
    for pszStep0007Path in objStep0007DivPaths:
        with open(pszStep0007Path, "r", encoding="utf-8", newline="") as objInputFile:
            objRows: List[List[str]] = list(csv.reader(objInputFile, delimiter="\t"))
        if objRows:
            while len(objRows[0]) < 2:
                objRows[0].append("")
            objRows[0][1] = "Div販管費"
        pszOutputBaseName: Optional[str] = build_step0008_div_output_file_name(os.path.basename(pszStep0007Path))
        if pszOutputBaseName is None:
            continue
        pszOutputPath: str = os.path.join(os.path.dirname(pszStep0007Path), pszOutputBaseName)
        with open(pszOutputPath, "w", encoding="utf-8", newline="") as objOutputFile:
            csv.writer(objOutputFile, delimiter="\t", lineterminator="\n").writerows(objRows)
        objOutputPaths.append(pszOutputPath)
    return objOutputPaths


def create_step0008_from_step0007_grp_paths(objStep0007GrpPaths: List[str]) -> List[str]:
    objOutputPaths: List[str] = []
    for pszStep0007Path in objStep0007GrpPaths:
        with open(pszStep0007Path, "r", encoding="utf-8", newline="") as objInputFile:
            objRows: List[List[str]] = list(csv.reader(objInputFile, delimiter="\t"))
        if objRows:
            while len(objRows[0]) < 2:
                objRows[0].append("")
            objRows[0][1] = "Div販管費"
        pszOutputBaseName: Optional[str] = build_step0008_grp_output_file_name(os.path.basename(pszStep0007Path))
        if pszOutputBaseName is None:
            continue
        pszOutputPath: str = os.path.join(os.path.dirname(pszStep0007Path), pszOutputBaseName)
        with open(pszOutputPath, "w", encoding="utf-8", newline="") as objOutputFile:
            csv.writer(objOutputFile, delimiter="\t", lineterminator="\n").writerows(objRows)
        objOutputPaths.append(pszOutputPath)
    return objOutputPaths


def main() -> int:
    objInputFiles: List[str] = sys.argv[1:]
    if not objInputFiles:
        print("Error: no input files specified.", file=sys.stderr)
        return 1

    objPlInputFiles: List[str] = []
    objStep0002InputFiles: List[str] = []
    objManhourInputFiles: List[str] = []
    for pszPath in objInputFiles:
        pszBaseName: str = os.path.basename(pszPath)
        if is_pl_tsv_file_name(pszBaseName):
            objPlInputFiles.append(pszPath)
        elif is_step0002_tsv_file_name(pszBaseName):
            objStep0002InputFiles.append(pszPath)
        elif is_manhour_tsv_file_name(pszBaseName):
            objManhourInputFiles.append(pszPath)

    if objStep0002InputFiles and objManhourInputFiles:
        objStep0002Map: Dict[str, str] = build_year_month_path_map(objStep0002InputFiles)
        objManhourMap: Dict[str, str] = build_year_month_path_map(objManhourInputFiles)
        try:
            objStep0003OutputPaths, objSkippedYearMonths = run_step0003_with_maps(
                objStep0002Map,
                objManhourMap,
            )
        except Exception as exc:  # noqa: BLE001
            print(f"Error: failed to process step0003. Detail = {exc}", file=sys.stderr)
            return 1

        print(f"Processed step0003 TSV count: {len(objStep0003OutputPaths)}")
        for pszOutputPath in objStep0003OutputPaths:
            print(f"Output(step0003): {pszOutputPath}")
        objStep0004Paths: List[str] = run_step0004_for_paths(objStep0003OutputPaths)
        print(f"Processed step0004 TSV count: {len(objStep0004Paths)}")
        for pszOutputPath in objStep0004Paths:
            print(f"Output(step0004): {pszOutputPath}")
        objPeriodPaths, objPeriodWarningPaths, objPeriodErrorPaths = aggregate_step0004_for_periods(objStep0004Paths)
        if objPeriodPaths:
            print(f"Processed step0004 period TSV count: {len(objPeriodPaths)}")
            for pszOutputPath in objPeriodPaths:
                print(f"Output(step0004_period): {pszOutputPath}")
            pszSummaryExcelPath: Optional[str] = create_step0004_summary_excel(objPeriodPaths)
            if pszSummaryExcelPath is not None:
                print(f"Output(step0004_summary_xlsx): {pszSummaryExcelPath}")
            objStep0005Paths: List[str] = create_step0005_from_period_paths(objPeriodPaths)
            print(f"Processed step0005 TSV count: {len(objStep0005Paths)}")
            for pszOutputPath in objStep0005Paths:
                print(f"Output(step0005): {pszOutputPath}")
            objStep0006DivPaths, objStep0006GrpPaths = create_step0006_from_step0005_paths(objStep0005Paths)
            print(f"Processed step0006 Div TSV count: {len(objStep0006DivPaths)}")
            for pszOutputPath in objStep0006DivPaths:
                print(f"Output(step0006_div): {pszOutputPath}")
            print(f"Processed step0006 Grp TSV count: {len(objStep0006GrpPaths)}")
            for pszOutputPath in objStep0006GrpPaths:
                print(f"Output(step0006_grp): {pszOutputPath}")
            objStep0007Paths: List[str] = create_step0007_from_step0006_div_paths(objStep0006DivPaths)
            print(f"Processed step0007 Div TSV count: {len(objStep0007Paths)}")
            for pszOutputPath in objStep0007Paths:
                print(f"Output(step0007_div): {pszOutputPath}")
            pszStep0007SummaryPath, pszStep0007SummaryErrorPath = create_step0007_div_summary_excel(objStep0007Paths)
            if pszStep0007SummaryPath is not None:
                print(f"Output(step0007_div_summary_xlsx): {pszStep0007SummaryPath}")
            if pszStep0007SummaryErrorPath is not None:
                print(f"WarningErrorFile: {pszStep0007SummaryErrorPath}")
            objStep0007GrpPaths: List[str] = create_step0007_from_step0006_grp_paths(objStep0006GrpPaths)
            print(f"Processed step0007 Grp TSV count: {len(objStep0007GrpPaths)}")
            for pszOutputPath in objStep0007GrpPaths:
                print(f"Output(step0007_grp): {pszOutputPath}")
            pszStep0007GrpSummaryPath, pszStep0007GrpSummaryErrorPath = create_step0007_grp_summary_excel(objStep0007GrpPaths)
            if pszStep0007GrpSummaryPath is not None:
                print(f"Output(step0007_grp_summary_xlsx): {pszStep0007GrpSummaryPath}")
            if pszStep0007GrpSummaryErrorPath is not None:
                print(f"WarningErrorFile: {pszStep0007GrpSummaryErrorPath}")
            objStep0008DivPaths: List[str] = create_step0008_from_step0007_div_paths(objStep0007Paths)
            print(f"Processed step0008 Div TSV count: {len(objStep0008DivPaths)}")
            for pszOutputPath in objStep0008DivPaths:
                print(f"Output(step0008_div): {pszOutputPath}")
            objStep0008GrpPaths: List[str] = create_step0008_from_step0007_grp_paths(objStep0007GrpPaths)
            print(f"Processed step0008 Grp TSV count: {len(objStep0008GrpPaths)}")
            for pszOutputPath in objStep0008GrpPaths:
                print(f"Output(step0008_grp): {pszOutputPath}")
        for pszWarningPath in objPeriodWarningPaths:
            print(f"WarningErrorFile: {pszWarningPath}")
        for pszErrorPath in objPeriodErrorPaths:
            print(f"ErrorFile: {pszErrorPath}")
        if objSkippedYearMonths:
            print(f"Skipped step0003 count: {len(objSkippedYearMonths)}")
            for pszYearMonth in objSkippedYearMonths:
                print(f"Skipped(step0003): {pszYearMonth}")
        return 0

    if not objPlInputFiles:
        print("Error: no PL TSV files matched the expected format.", file=sys.stderr)
        return 1

    objStep0001OutputPaths: List[str] = []
    for pszPlPath in objPlInputFiles:
        if not os.path.isfile(pszPlPath):
            print(f"Error: input file not found: {pszPlPath}", file=sys.stderr)
            return 1
        try:
            pszStep0001Path: str = process_one_pl_tsv_to_step0001(pszPlPath)
        except Exception as exc:  # noqa: BLE001
            print(f"Error: failed to process {pszPlPath}. Detail = {exc}", file=sys.stderr)
            return 1
        objStep0001OutputPaths.append(pszStep0001Path)

    objStep0002OutputPaths: List[str] = []
    objSkippedStep0001Paths: List[str] = []
    objErrorPaths: List[str] = []
    for pszStep0001Path in objStep0001OutputPaths:
        try:
            objStep0002Result = process_one_step0001_tsv_to_step0002(pszStep0001Path)
        except Exception as exc:  # noqa: BLE001
            print(f"Error: failed to process step0002 {pszStep0001Path}. Detail = {exc}", file=sys.stderr)
            return 1
        pszStep0002Path, pszErrorPath = objStep0002Result
        if pszStep0002Path is None:
            objSkippedStep0001Paths.append(pszStep0001Path)
            continue
        objStep0002OutputPaths.append(pszStep0002Path)
        if pszErrorPath is not None:
            objErrorPaths.append(pszErrorPath)

    print(f"Processed PL TSV count: {len(objStep0001OutputPaths)}")
    for pszOutputPath in objStep0001OutputPaths:
        print(f"Output(step0001): {pszOutputPath}")

    print(f"Processed step0002 TSV count: {len(objStep0002OutputPaths)}")
    for pszOutputPath in objStep0002OutputPaths:
        print(f"Output(step0002): {pszOutputPath}")

    if objSkippedStep0001Paths:
        print(f"Skipped step0002 count: {len(objSkippedStep0001Paths)}")
        for pszPath in objSkippedStep0001Paths:
            print(f"Skipped(step0002): {pszPath}")

    if objErrorPaths:
        print(f"Warning error file count: {len(objErrorPaths)}")
        for pszPath in objErrorPaths:
            print(f"WarningErrorFile: {pszPath}")

    if objManhourInputFiles:
        objStep0002Map = build_year_month_path_map(objStep0002OutputPaths)
        objManhourMap = build_year_month_path_map(objManhourInputFiles)
        try:
            objStep0003OutputPaths, objSkippedYearMonths = run_step0003_with_maps(
                objStep0002Map,
                objManhourMap,
            )
        except Exception as exc:  # noqa: BLE001
            print(f"Error: failed to process step0003. Detail = {exc}", file=sys.stderr)
            return 1
        print(f"Processed step0003 TSV count: {len(objStep0003OutputPaths)}")
        for pszOutputPath in objStep0003OutputPaths:
            print(f"Output(step0003): {pszOutputPath}")
        objStep0004Paths: List[str] = run_step0004_for_paths(objStep0003OutputPaths)
        print(f"Processed step0004 TSV count: {len(objStep0004Paths)}")
        for pszOutputPath in objStep0004Paths:
            print(f"Output(step0004): {pszOutputPath}")
        objPeriodPaths, objPeriodWarningPaths, objPeriodErrorPaths = aggregate_step0004_for_periods(objStep0004Paths)
        if objPeriodPaths:
            print(f"Processed step0004 period TSV count: {len(objPeriodPaths)}")
            for pszOutputPath in objPeriodPaths:
                print(f"Output(step0004_period): {pszOutputPath}")
            pszSummaryExcelPath = create_step0004_summary_excel(objPeriodPaths)
            if pszSummaryExcelPath is not None:
                print(f"Output(step0004_summary_xlsx): {pszSummaryExcelPath}")
            objStep0005Paths = create_step0005_from_period_paths(objPeriodPaths)
            print(f"Processed step0005 TSV count: {len(objStep0005Paths)}")
            for pszOutputPath in objStep0005Paths:
                print(f"Output(step0005): {pszOutputPath}")
            objStep0006DivPaths, objStep0006GrpPaths = create_step0006_from_step0005_paths(objStep0005Paths)
            print(f"Processed step0006 Div TSV count: {len(objStep0006DivPaths)}")
            for pszOutputPath in objStep0006DivPaths:
                print(f"Output(step0006_div): {pszOutputPath}")
            print(f"Processed step0006 Grp TSV count: {len(objStep0006GrpPaths)}")
            for pszOutputPath in objStep0006GrpPaths:
                print(f"Output(step0006_grp): {pszOutputPath}")
            objStep0007Paths: List[str] = create_step0007_from_step0006_div_paths(objStep0006DivPaths)
            print(f"Processed step0007 Div TSV count: {len(objStep0007Paths)}")
            for pszOutputPath in objStep0007Paths:
                print(f"Output(step0007_div): {pszOutputPath}")
            pszStep0007SummaryPath, pszStep0007SummaryErrorPath = create_step0007_div_summary_excel(objStep0007Paths)
            if pszStep0007SummaryPath is not None:
                print(f"Output(step0007_div_summary_xlsx): {pszStep0007SummaryPath}")
            if pszStep0007SummaryErrorPath is not None:
                print(f"WarningErrorFile: {pszStep0007SummaryErrorPath}")
            objStep0007GrpPaths: List[str] = create_step0007_from_step0006_grp_paths(objStep0006GrpPaths)
            print(f"Processed step0007 Grp TSV count: {len(objStep0007GrpPaths)}")
            for pszOutputPath in objStep0007GrpPaths:
                print(f"Output(step0007_grp): {pszOutputPath}")
            pszStep0007GrpSummaryPath, pszStep0007GrpSummaryErrorPath = create_step0007_grp_summary_excel(objStep0007GrpPaths)
            if pszStep0007GrpSummaryPath is not None:
                print(f"Output(step0007_grp_summary_xlsx): {pszStep0007GrpSummaryPath}")
            if pszStep0007GrpSummaryErrorPath is not None:
                print(f"WarningErrorFile: {pszStep0007GrpSummaryErrorPath}")
            objStep0008DivPaths: List[str] = create_step0008_from_step0007_div_paths(objStep0007Paths)
            print(f"Processed step0008 Div TSV count: {len(objStep0008DivPaths)}")
            for pszOutputPath in objStep0008DivPaths:
                print(f"Output(step0008_div): {pszOutputPath}")
            objStep0008GrpPaths: List[str] = create_step0008_from_step0007_grp_paths(objStep0007GrpPaths)
            print(f"Processed step0008 Grp TSV count: {len(objStep0008GrpPaths)}")
            for pszOutputPath in objStep0008GrpPaths:
                print(f"Output(step0008_grp): {pszOutputPath}")
        for pszWarningPath in objPeriodWarningPaths:
            print(f"WarningErrorFile: {pszWarningPath}")
        for pszErrorPath in objPeriodErrorPaths:
            print(f"ErrorFile: {pszErrorPath}")
        if objSkippedYearMonths:
            print(f"Skipped step0003 count: {len(objSkippedYearMonths)}")
            for pszYearMonth in objSkippedYearMonths:
                print(f"Skipped(step0003): {pszYearMonth}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
